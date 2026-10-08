"""The files that go with the results page: MNE's own reports, the
short summary document (Word, and PDF through LibreOffice when it is
installed) and a zip of the folder to send."""
from __future__ import annotations

import shutil
import subprocess
import zipfile
from pathlib import Path

import numpy as np

import mne

from . import pipeline as P
from .analyses import SRT_LABELS


def mne_reports(blocks, cleaned, results, out: Path) -> list[str]:
    names = []
    for b in blocks:
        c = cleaned[b.mode]
        rep = mne.Report(title=f"{b.mode.upper()}: {b.bdf.name}", verbose="ERROR")
        rep.add_events(b.eeg_events, title="Markers on the Status channel",
                       sfreq=b.raw.info["sfreq"])
        rep.add_raw(c.raw, title="Cleaned recording (0.1-40 Hz, average reference, ICA)",
                    psd=True, butterfly=False)
        rep.add_ica(c.ica, title="ICA: eye components removed", inst=c.raw_ica_fit,
                    picks=c.ica.exclude, n_jobs=1)
        if b.mode == "srt" and "srt" in results:
            r = results["srt"]
            t = r.trials
            ep = P.make_epochs(c.raw, t.flash_sample, np.full(len(t), 30), t,
                               -0.2, 0.6, (-0.2, 0.0))
            rep.add_epochs(ep, title="Flash epochs (all trials)", psd=False)
            evs = [r.stim[k] for k in ["all_correct", *SRT_LABELS] if k in r.stim]
            titles = ["All correct flashes"] + [SRT_LABELS[k] for k in SRT_LABELS if k in r.stim]
            rep.add_evokeds(evs, titles=[f"Flash: {x}" for x in titles], n_time_points=7)
            evs = [r.resp[k] for k in ("correct", "error", "anticipation") if k in r.resp]
            rep.add_evokeds(evs, titles=[f"Press (force onset): {k}" for k in ("correct", "error", "anticipation") if k in r.resp],
                            n_time_points=7)
        if b.mode == "buzz_hunt" and "buzz_hunt" in results:
            r = results["buzz_hunt"]
            evs = [r.erp[k] for k in ("localisation", "gap", "span") if k in r.erp]
            rep.add_evokeds(evs, titles=[f"Buzz: {k}" for k in ("localisation", "gap", "span") if k in r.erp],
                            n_time_points=7)
        name = f"mne_report_{b.mode}.html"
        rep.save(out / name, open_browser=False, overwrite=True, verbose="ERROR")
        names.append(name)
    return names


def summary_doc(s: dict, figs: Path, out: Path, methods_short: list[str],
                limits: list[str], refs: list[str]) -> list[str]:
    from docx import Document
    from docx.shared import Cm, Pt
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.top_margin = sec.bottom_margin = Cm(1.8)
    from . import findings as FD
    srt, bz = s.get("srt", {}), s.get("buzz", {})
    fd = FD.build(s)
    doc.add_heading(f"EEG results: {s.get('participant')}, {s.get('date_long')}", 0)
    doc.add_paragraph(f"{s.get('participant')}, {s.get('age')}, {s.get('dominant_hand')}-handed. Lab serial reaction "
                      f"time task (SRT, {srt.get('n_trials')} trials) and Buzz Hunt ({bz.get('trials')} trials) on the "
                      f"finger device, recorded with the lab's BioSemi ActiveTwo (64 channels, 512 Hz). Analysed with "
                      f"MNE-Python {s['software']['mne']}. The full interactive results page (EEG_results.html) and "
                      f"MNE's reports are in the same folder.")
    doc.add_heading("Key findings", 1)
    for i in fd["items"]:
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(i["head"] + " ").bold = True
        para.add_run(i["text"])
    doc.add_paragraph("In plain words: " + fd["plain"])
    doc.add_paragraph("One person, one session: every p value compares this person's trials with each other. "
                      "A finding is claimed only when its primary test passes; the N2 and P3 are Holm-corrected "
                      "as a pair, and every other comparison is exploratory.").runs[0].italic = True
    for name, cap in (("10_srt_behaviour.png", "SRT: median RT and accuracy per block."),
                      ("12_srt_erp_conditions.png", "SRT: response to the flash by phase."),
                      ("13_srt_erp_topomaps.png", "SRT: N2 and P3 scalp maps; right, random minus learned."),
                      ("19_srt_overlap_corrected.png", "SRT: regression ERPs with overlap removed."),
                      ("18_srt_rhythms_by_block.png", "SRT: band power block by block."),
                      ("21_buzz_erp.png", "Buzz Hunt: touch response."),
                      ("22_buzz_time_frequency.png", "Buzz Hunt: rhythm change after the buzz."),
                      ("01_markers.png", "Marker timing against the game's clock.")):
        f = figs / name
        if f.is_file():
            doc.add_picture(str(f), width=Cm(17))
            doc.add_paragraph(cap).runs[0].italic = True
    doc.add_heading("Methods in brief", 1)
    for x in methods_short:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("Limits", 1)
    for x in limits:
        doc.add_paragraph(x, style="List Bullet")
    doc.add_heading("References", 1)
    for x in refs:
        doc.add_paragraph(x)
    path = out / "EEG_results_summary.docx"
    doc.save(path)
    made = [path.name]
    soffice = shutil.which("soffice")
    if soffice:
        subprocess.run([soffice, "-env:UserInstallation=file:///private/tmp/claude-501/lo_profile",
                        "--headless", "--convert-to", "pdf", "--outdir", str(out), str(path)],
                       check=False, capture_output=True, timeout=300)
        if (out / "EEG_results_summary.pdf").is_file():
            made.append("EEG_results_summary.pdf")
    return made


def zip_folder(out: Path, name: str) -> Path:
    z = out.parent / f"{name}.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in sorted(out.rglob("*")):
            if f.is_file() and f.name != "cache.pkl" and not f.name.startswith("."):
                zf.write(f, Path(name) / f.relative_to(out))
    return z
